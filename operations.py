"""Phase 2: transactional imports, machine-readable labels and stock approvals.
All endpoints use the same authentication / role scoping as the core application.
"""
import base64
import csv
import hashlib
import io
import json
import re
import secrets
from client_types import tracking_for, tracking_text, check_available
import unicodedata
from datetime import datetime, timedelta
from zipfile import ZipFile, BadZipFile
from flask import request, jsonify, Response
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
import qrcode
from barcode import Code128
from barcode.writer import SVGWriter

COLUMNS = ['reference_externe', 'destinataire', 'telephone', 'adresse', 'ville', 'montant', 'produit', 'note']
MAX_ROWS = 500
MAX_FILE = 2 * 1024 * 1024


def normalize(value):
    value = unicodedata.normalize('NFKD', str(value or '').strip().casefold())
    return ''.join(c for c in value if not unicodedata.combining(c))


def csv_safe(value):
    s = str(value if value is not None else '')
    return "'" + s if s.lstrip().startswith(('=', '+', '-', '@')) else s


def codes_for(tracking):
    qr = qrcode.QRCode(error_correction=qrcode.constants.ERROR_CORRECT_M, box_size=7, border=4)
    qr.add_data(tracking)
    qr.make(fit=True)
    b = io.BytesIO()
    qr.make_image(fill_color='black', back_color='white').save(b, format='PNG')
    barcode = Code128(tracking, writer=SVGWriter())
    svg = barcode.render({'module_width': 0.32, 'module_height': 13, 'quiet_zone': 3.5,
                          'font_size': 9, 'text_distance': 4, 'write_text': True})
    return {'qr': 'data:image/png;base64,' + base64.b64encode(b.getvalue()).decode(),
            'barcode': 'data:image/svg+xml;base64,' + base64.b64encode(svg).decode()}


def read_import(raw, filename, Error):
    """Read a bounded sheet. Formulas and numerical telephone cells are rejected later."""
    if not raw or len(raw) > MAX_FILE:
        raise Error('Fichier vide ou supérieur à 2 Mo.')
    suffix = filename.lower().rsplit('.', 1)[-1]
    if suffix == 'xlsx':
        try:
            with ZipFile(io.BytesIO(raw)) as z:
                if len(z.infolist()) > 500 or sum(i.file_size for i in z.infolist()) > 20 * 1024 * 1024:
                    raise Error('Classeur trop volumineux après décompression (maximum 20 Mo).')
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
            try:
                ws = wb['Colis'] if 'Colis' in wb.sheetnames else wb.worksheets[0]
                # Do not trust the worksheet dimension element supplied by the file.
                if ws.max_column and ws.max_column > 20:
                    raise Error('Le classeur contient trop de colonnes.')
                ws.reset_dimensions()
                values = []
                for cells in ws.iter_rows(max_col=20):
                    if len(values) >= MAX_ROWS + 1:
                        raise Error(f'Maximum {MAX_ROWS} lignes par import, sans lignes vides supplémentaires formatées.')
                    row = [cell.value for cell in cells]
                    # Boolean/date/etc. cells are not silently coerced to valid domain fields.
                    for i, cell in enumerate(cells):
                        if cell.data_type == 'f':
                            row[i] = str(cell.value)
                    values.append(row)
            finally:
                wb.close()
        except Error:
            raise
        except Exception:
            raise Error('Classeur illisible. Utilisez le modèle .xlsx, sans mot de passe ni macros.')
    elif suffix == 'csv':
        try:
            content = raw.decode('utf-8-sig')
            try:
                dialect = csv.Sniffer().sniff(content[:8192], delimiters=';,\t')
                reader = csv.reader(io.StringIO(content), dialect)
            except csv.Error:
                reader = csv.reader(io.StringIO(content), delimiter=';')
            values = []
            for row in reader:
                if len(row) > 20:
                    raise Error('Le fichier contient trop de colonnes.')
                values.append(row)
                if len(values) > MAX_ROWS + 1:
                    raise Error(f'Maximum {MAX_ROWS} lignes par import.')
        except UnicodeDecodeError:
            raise Error('Le CSV doit être encodé en UTF-8. Vous pouvez utiliser le modèle Excel.')
        except csv.Error:
            raise Error('CSV illisible. Vérifiez les séparateurs et les guillemets.')
    else:
        raise Error('Formats acceptés : .xlsx et .csv uniquement.')
    if not values:
        raise Error('Le fichier ne contient aucune ligne.')
    header = [normalize(x).replace(' ', '_') for x in values[0]]
    while header and not header[-1]:
        header.pop()
    if len(header) != len(set(header)) or any(not x for x in header):
        raise Error('Les en-têtes doivent être uniques et sans colonne vide.')
    missing = [x for x in COLUMNS[:6] if x not in header]
    if missing:
        raise Error('Colonnes obligatoires manquantes : ' + ', '.join(missing))
    unknown = [x for x in header if x not in COLUMNS]
    if unknown:
        raise Error('Colonnes inconnues : ' + ', '.join(unknown) + '. Téléchargez le modèle à jour.')
    records = []
    for num, row in enumerate(values[1:], 2):
        if not any(x is not None and str(x).strip() for x in row):
            continue
        if num > MAX_ROWS + 1:
            raise Error(f'Maximum {MAX_ROWS} lignes par import.')
        if any(x is not None and str(x).strip() for x in row[len(header):]):
            raise Error(f'Ligne {num} : une valeur ne correspond à aucune colonne.')
        records.append((num, {key: row[i] if i < len(row) else '' for i, key in enumerate(header)}))
    if not records:
        raise Error('Ajoutez au moins une ligne de colis après les en-têtes.')
    return records


def register_operations(app, services):
    conn, auth, user, scope = (services[k] for k in ('conn', 'auth', 'user', 'scope'))
    parcel, event, text, number, now, Error = (services[k] for k in ('parcel', 'event', 'text', 'number', 'now', 'APIError'))
    with conn() as c:
        c.executescript('''
        CREATE TABLE IF NOT EXISTS import_batches (
            id TEXT PRIMARY KEY, user_id INTEGER NOT NULL REFERENCES users(id),
            client_id INTEGER NOT NULL REFERENCES users(id), filename TEXT NOT NULL,
            digest TEXT NOT NULL, payload TEXT NOT NULL, status TEXT NOT NULL DEFAULT 'Prêt',
            created_at TEXT NOT NULL, committed_at TEXT, result TEXT);
        CREATE TABLE IF NOT EXISTS parcel_external_refs (
            client_id INTEGER NOT NULL REFERENCES users(id), reference TEXT NOT NULL,
            parcel_id INTEGER UNIQUE NOT NULL REFERENCES parcels(id),
            PRIMARY KEY(client_id,reference));
        CREATE TABLE IF NOT EXISTS stock_requests (
            id INTEGER PRIMARY KEY, product_id INTEGER NOT NULL REFERENCES products(id),
            client_id INTEGER NOT NULL REFERENCES users(id), created_by INTEGER NOT NULL REFERENCES users(id),
            kind TEXT NOT NULL CHECK(kind IN ('Entrée','Sortie')), quantity INTEGER NOT NULL CHECK(quantity>0),
            reason TEXT NOT NULL, status TEXT NOT NULL DEFAULT 'En attente',
            admin_note TEXT NOT NULL DEFAULT '', processed_by INTEGER REFERENCES users(id),
            created_at TEXT NOT NULL, processed_at TEXT,
            movement_id INTEGER UNIQUE REFERENCES movements(id));
        CREATE INDEX IF NOT EXISTS idx_import_owner ON import_batches(user_id,created_at);
        CREATE INDEX IF NOT EXISTS idx_stock_requests_client ON stock_requests(client_id,status);
        CREATE INDEX IF NOT EXISTS idx_parcels_tracking ON parcels(tracking);
        ''')

    @app.errorhandler(413)
    def import_too_large(e):
        return jsonify(error='Fichier trop volumineux : maximum 2 Mo (envoyez un fichier légèrement inférieur à cette limite).'), 413

    def client_for(c, u, data):
        cid = u['id'] if u['role'] == 'client' else data.get('client_id')
        cl = c.execute("SELECT id,name,company,client_type FROM users WHERE id=? AND role='client' AND active=1", (cid,)).fetchone()
        if not cl:
            raise Error('Choisissez un client actif pour cet import.')
        return dict(cl)

    def validate_record(c, rownum, raw, cities, cid, seen, client_type):
        errors = []
        clean = {}
        for k in COLUMNS:
            v = raw.get(k)
            s = '' if v is None else str(v).strip()
            clean[k] = s
            if len(s) > (80 if k == 'reference_externe' else 300):
                errors.append(f'{k} : texte trop long.')
            if s.startswith('='):
                errors.append(f'{k} : les formules ne sont pas acceptées.')
        for k in COLUMNS[:6]:
            if not clean[k]:
                errors.append(f'{k} : champ obligatoire.')
        ref = clean['reference_externe']
        refkey = ref.casefold()
        if ref and refkey in seen:
            errors.append('Référence externe répétée dans ce fichier.')
        seen.add(refkey)
        if ref and c.execute('SELECT 1 FROM parcel_external_refs WHERE client_id=? AND reference=?', (cid, refkey)).fetchone():
            errors.append('Référence externe déjà importée pour ce client.')
        if client_type == 'societe_livraison':
            try:
                if not isinstance(raw.get('reference_externe'),str):raise Error('Tracking : utilisez une cellule texte, pas un nombre (zéros initiaux).')
                tracking_for(c,cid,ref,Error,generate=False)
            except Error as e:errors.append(e.msg)
        telephone = re.sub(r'[\s().-]', '', clean['telephone'])
        if isinstance(raw.get('telephone'), (int, float)) or not re.fullmatch(r'(?:0[5-8]\d{8}|\+212[5-8]\d{8})', telephone):
            errors.append('Téléphone marocain invalide : format texte 06… / 07… ou +212… (conservez le zéro).')
        clean['telephone'] = telephone
        city = cities.get(normalize(clean['ville']))
        if not city:
            errors.append('Ville inconnue ou fermée à la livraison. Consultez la feuille Villes / les paramètres.')
        try:
            amount_text = clean['montant'].replace('\u202f', '').replace('\xa0', '').replace(' ', '').replace(',', '.')
            amount = number({'montant': amount_text}, 'montant')
        except Error:
            amount = None
            errors.append('Montant invalide : nombre entre 0 et 1 000 000 MAD.')
        return {'line': rownum, 'reference': ref, 'recipient': clean['destinataire'], 'phone': telephone,
                'address': clean['adresse'], 'city': city['name'] if city else clean['ville'],
                'city_id': city['id'] if city else None, 'amount': amount,
                'fee': city['fee'] if city else None, 'return_fee': city['return_fee'] if city else None,
                'product': clean['produit'], 'note': clean['note'], 'errors': errors,
                'tracking': ref if client_type == 'societe_livraison' else None}

    @app.get('/api/imports/template.<fmt>')
    @auth('admin', 'client')
    def import_template(fmt):
        with conn() as c:
            cities = [dict(r) for r in c.execute('SELECT * FROM cities WHERE delivery=1 ORDER BY name')]
        example = ['EXEMPLE-001', 'Destinataire exemple', '0600000000', 'Adresse exemple à remplacer', cities[0]['name'] if cities else '', 199, 'Produit exemple', 'Ligne de démonstration à remplacer']
        if fmt == 'csv':
            b = io.StringIO(); w = csv.writer(b, delimiter=';'); w.writerow(COLUMNS); w.writerow([csv_safe(v) for v in example])
            return Response('\ufeff'+b.getvalue(), mimetype='text/csv', headers={'Content-Disposition': 'attachment; filename="ORIENTAL24-modele-colis.csv"'})
        if fmt != 'xlsx':
            raise Error('Format inconnu.', 404)
        wb = Workbook(); ws = wb.active; ws.title = 'Colis'
        ws.append(COLUMNS); ws.append(example); ws.freeze_panes = 'A2'
        for cell in ws[2]:
            if isinstance(cell.value, str): cell.data_type = 's'
        ws.auto_filter.ref = 'A1:H2'
        for i, width in enumerate([24, 26, 22, 40, 24, 16, 26, 45], 1):
            ws.column_dimensions[get_column_letter(i)].width = width
        for cell in ws[1]:
            cell.font = Font(color='FFFFFF', bold=True)
            cell.fill = PatternFill('solid', fgColor='091F3F')
        for row in ws.iter_rows(min_row=2, max_row=MAX_ROWS+1):
            row[0].number_format = '@'; row[2].number_format = '@'; row[5].number_format = '0.00'
        cities_ws = wb.create_sheet('Villes'); cities_ws.append(['Ville de livraison ouverte', 'Région', 'Frais livraison MAD', 'Frais retour MAD'])
        for city in cities:
            # Explicit text cell types prevent a city name from becoming an Excel formula.
            cities_ws.append([city['name'], city['region'], city['fee'], city['return_fee']])
            for cell in cities_ws[cities_ws.max_row][:2]: cell.data_type = 's'
        for col in 'ABCD': cities_ws.column_dimensions[col].width = 28
        if cities:
            dv = DataValidation(type='list', formula1=f"'Villes'!$A$2:$A${len(cities)+1}", allow_blank=False)
            dv.error = 'Choisissez une ville de la feuille Villes.'; dv.showErrorMessage = True
            ws.add_data_validation(dv); dv.add(f'E2:E{MAX_ROWS+1}')
        guide = wb.create_sheet('Guide')
        for value in [
            'ORIENTAL24 — Modèle d’import des colis',
            'Remplacez la ligne EXEMPLE-001 par vos propres données. Elle sera importée si vous la conservez.',
            'Maximum 500 lignes, fichier .xlsx ou CSV UTF-8, taille inférieure à 2 Mo.',
            'Vendeur : reference_externe reste votre référence de commande ; ORIENTAL24 génère le tracking.',
            'Société de livraison : reference_externe est votre tracking, conservé tel quel, y compris sa casse et ses zéros. Cellule texte obligatoire.',
            'Tracking société : 1 à 80 caractères A–Z/a–z, chiffres, point, tiret, underscore, / ; début alphanumérique. Unique sur la plateforme, sans distinction de casse.',
            'telephone : cellule texte, format 06/07… ou +212…, conservez le zéro initial.',
            'montant : somme à collecter, sans MAD. Les décimales avec point ou virgule sont acceptées.',
            'ville : nom exact d’une ville active. Les accents et la casse sont tolérés.',
            'Les tarifs sont contrôlés à la confirmation. Aucun colis n’est créé au stade de prévisualisation.',
            'Si une ligne contient une erreur, corrigez votre fichier et recommencez : aucun import partiel.',
            'La prévisualisation expire après 30 minutes. Aucune formule Excel ni macro n’est exécutée.'
        ]: guide.append([value])
        guide.column_dimensions['A'].width = 115
        for row in guide:
            row[0].alignment = Alignment(wrap_text=True, vertical='top')
            guide.row_dimensions[row[0].row].height = 32
        b = io.BytesIO(); wb.save(b); wb.close()
        return Response(b.getvalue(), mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={'Content-Disposition': 'attachment; filename="ORIENTAL24-modele-colis.xlsx"'})

    @app.post('/api/imports/preview')
    @auth('admin', 'client')
    def preview_import():
        u = user()
        upload = request.files.get('file')
        if not upload or not upload.filename:
            raise Error('Sélectionnez un fichier Excel ou CSV.')
        filename = upload.filename.replace('\\', '/').split('/')[-1][:180]
        raw = upload.read(MAX_FILE+1)
        parsed = read_import(raw, filename, Error)
        with conn() as c:
            client = client_for(c, u, request.form)
            cities = {normalize(r['name']): dict(r) for r in c.execute('SELECT * FROM cities WHERE delivery=1')}
            seen = set()
            rows = [validate_record(c, n, r, cities, client['id'], seen, client['client_type']) for n, r in parsed]
            invalid = sum(bool(r['errors']) for r in rows)
            payload = {'rows': rows, 'total': len(rows), 'valid': len(rows)-invalid, 'invalid': invalid,
                       'client': client['company'] or client['name'], 'client_type': client['client_type']}
            token = secrets.token_hex(16)
            status = 'À corriger' if invalid else 'Prêt'
            c.execute('INSERT INTO import_batches(id,user_id,client_id,filename,digest,payload,status,created_at) VALUES(?,?,?,?,?,?,?,?)',
                      (token, u['id'], client['id'], filename, hashlib.sha256(raw).hexdigest(), json.dumps(payload, ensure_ascii=False), status, now()))
        return jsonify(id=token, filename=filename, status=status, created_at=now(), **payload)

    def get_batch(c, token, uid):
        b = c.execute('SELECT * FROM import_batches WHERE id=? AND user_id=?', (token, uid)).fetchone()
        if not b:
            raise Error('Import introuvable.', 404)
        return b

    @app.get('/api/imports')
    @auth('admin', 'client')
    def list_imports():
        with conn() as c:
            rows = c.execute('SELECT * FROM import_batches WHERE user_id=? ORDER BY created_at DESC,rowid DESC LIMIT 20', (user()['id'],)).fetchall()
            return jsonify([{'id': b['id'], 'filename': b['filename'], 'status': b['status'], 'created_at': b['created_at'],
                             'committed_at': b['committed_at'], **{k: json.loads(b['payload'])[k] for k in ('total', 'valid', 'invalid', 'client')}} for b in rows])

    @app.get('/api/imports/<token>')
    @auth('admin', 'client')
    def import_detail(token):
        with conn() as c:
            b = get_batch(c, token, user()['id'])
            return jsonify(id=b['id'], filename=b['filename'], status=b['status'], created_at=b['created_at'],
                           committed_at=b['committed_at'], result=json.loads(b['result']) if b['result'] else None,
                           **json.loads(b['payload']))

    @app.get('/api/imports/<token>/errors')
    @auth('admin', 'client')
    def import_errors(token):
        with conn() as c:
            b = get_batch(c, token, user()['id'])
            payload = json.loads(b['payload'])
        output = io.StringIO(); writer = csv.writer(output, delimiter=';')
        writer.writerow(['Ligne', 'Référence externe', 'Destinataire', 'Erreurs'])
        for row in payload['rows']:
            if row['errors']:
                writer.writerow([row['line'], csv_safe(row['reference']), csv_safe(row['recipient']), csv_safe(' | '.join(row['errors']))])
        return Response('\ufeff'+output.getvalue(), mimetype='text/csv', headers={'Content-Disposition': 'attachment; filename="ORIENTAL24-erreurs-import.csv"'})

    @app.post('/api/imports/<token>/commit')
    @auth('admin', 'client')
    def commit_import(token):
        u = user()
        with conn() as c:
            c.execute('BEGIN IMMEDIATE')
            b = get_batch(c, token, u['id'])
            if b['status'] == 'Importé':
                return jsonify(ok=True, already_imported=True, **json.loads(b['result']))
            payload = json.loads(b['payload'])
            if b['status'] != 'Prêt' or payload['invalid']:
                raise Error('Corrigez toutes les lignes et prévisualisez à nouveau le fichier.')
            if datetime.fromisoformat(b['created_at']) < datetime.now()-timedelta(minutes=30):
                raise Error('Cette prévisualisation a expiré. Importez à nouveau votre fichier.', 409)
            client = client_for(c, u, {'client_id': b['client_id']})
            if client['client_type'] != payload.get('client_type','vendeur'):
                raise Error('Le type du client a changé. Refaites la prévisualisation.',409)
            created = []
            for r in payload['rows']:
                city = c.execute('SELECT * FROM cities WHERE id=? AND delivery=1', (r['city_id'],)).fetchone()
                if not city or city['name'] != r['city'] or city['fee'] != r['fee'] or city['return_fee'] != r['return_fee']:
                    raise Error(f"La couverture ou les tarifs de {r['city']} ont changé. Refaites la prévisualisation.", 409)
                if c.execute('SELECT 1 FROM parcel_external_refs WHERE client_id=? AND reference=?', (b['client_id'], r['reference'].casefold())).fetchone():
                    raise Error(f"La référence {r['reference']} a déjà été importée. Aucun colis ajouté.", 409)
                tracking = tracking_for(c,b['client_id'],r['reference'] if client['client_type']=='societe_livraison' else None,Error); t = now()
                fee, ret = services['tariff_fees_for'](c, b['client_id'], dict(city)) if services.get('tariff_fees_for') else (city['fee'], city['return_fee'])
                cur = c.execute('INSERT INTO parcels(tracking,client_id,recipient,phone,address,city_id,amount,fee,return_fee,product,note,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)',
                                (tracking, b['client_id'], r['recipient'], r['phone'], r['address'], city['id'], r['amount'], fee, ret, r['product'], r['note'], t, t))
                c.execute('INSERT INTO parcel_external_refs(client_id,reference,parcel_id) VALUES(?,?,?)', (b['client_id'], r['reference'].casefold(), cur.lastrowid))
                event(c, cur.lastrowid, 'Créé', f"Import {filename_display(b['filename'])} · Commande {r['reference']}", u)
                created.append({'id': cur.lastrowid, 'tracking': tracking, 'reference': r['reference']})
            result = {'count': len(created), 'parcels': created}
            c.execute("UPDATE import_batches SET status='Importé',committed_at=?,result=? WHERE id=?", (now(), json.dumps(result), token))
            return jsonify(ok=True, already_imported=False, **result)

    @app.post('/api/labels')
    @auth()
    def labels():
        data = request.get_json() or {}; ids = data.get('ids')
        if not isinstance(ids, list) or not 1 <= len(ids) <= 50 or any(type(i) is not int for i in ids):
            raise Error('Choisissez de 1 à 50 colis à imprimer.')
        u = user(); result = []
        with conn() as c:
            for pid in dict.fromkeys(ids):
                p = parcel(c, pid, u)
                extras = c.execute('SELECT ci.name city,u.name client,u.company company FROM cities ci JOIN users u ON u.id=? WHERE ci.id=?', (p['client_id'], p['city_id'])).fetchone()
                result.append({**p, **dict(extras), **codes_for(p['tracking'])})
        return jsonify(result)

    @app.post('/api/scan')
    @auth()
    def lookup_scan():
        payload = request.get_json() or {}
        tracking = tracking_text(payload.get('tracking'), Error)
        cond, args = scope(user())
        with conn() as c:
            rows = c.execute(f'SELECT p.id,p.tracking,p.client_id,u.company FROM parcels p JOIN users u ON u.id=p.client_id WHERE p.tracking=? COLLATE NOCASE AND {cond} ORDER BY p.id DESC LIMIT 12', [tracking]+args).fetchall()
            if not rows:
                raise Error('Aucun colis accessible avec cette référence.', 404)
            p = rows[0]
            if len(rows) > 1:
                chosen = payload.get('client_id')
                match = [r for r in rows if type(chosen) is int and r['client_id'] == chosen]
                if len(match) != 1:
                    companies = ', '.join(sorted({r['company'] or '—' for r in rows[:6]}))
                    raise Error('Référence identique chez plusieurs sociétés (' + companies + ') : ouvrez le bon colis via la recherche détaillée avant de scanner.', 409)
                p = match[0]
            return jsonify(dict(p))

    def find_product(c, pid, u):
        p = c.execute('SELECT * FROM products WHERE id=?', (pid,)).fetchone()
        if not p or (u['role'] == 'client' and p['client_id'] != u['id']):
            raise Error('Produit introuvable.', 404)
        return p

    @app.route('/api/stock/requests', methods=['GET', 'POST'])
    @auth('admin', 'client')
    def stock_requests():
        u = user()
        with conn() as c:
            if request.method == 'POST':
                d = request.get_json() or {}
                p = find_product(c, d.get('product_id'), u)
                kind = d.get('kind')
                if kind not in ('Entrée', 'Sortie'):
                    raise Error('Choisissez Entrée ou Sortie.')
                quantity = number(d, 'quantity', 1, 100000, True)
                if kind == 'Sortie' and quantity > p['quantity'] - services['reserved_stock'](c,p['id']):
                    raise Error('La quantité demandée dépasse le stock disponible.')
                cur = c.execute('INSERT INTO stock_requests(product_id,client_id,created_by,kind,quantity,reason,created_at) VALUES(?,?,?,?,?,?,?)',
                                (p['id'], p['client_id'], u['id'], kind, quantity, text(d, 'reason'), now()))
                return jsonify(ok=True, id=cur.lastrowid)
            cond, args = ('1=1', []) if u['role'] == 'admin' else ('r.client_id=?', [u['id']])
            rows = c.execute(f'''SELECT r.*,p.name product,p.reference sku,p.quantity available,
                cl.company company,cl.name client,creator.name creator,actor.name processor
                FROM stock_requests r JOIN products p ON p.id=r.product_id
                JOIN users cl ON cl.id=r.client_id JOIN users creator ON creator.id=r.created_by
                LEFT JOIN users actor ON actor.id=r.processed_by WHERE {cond} ORDER BY r.id DESC''', args)
            return jsonify([dict(r) for r in rows])

    @app.patch('/api/stock/requests/<int:rid>')
    @auth('admin', 'client')
    def decide_stock_request(rid):
        u = user(); d = request.get_json() or {}; status = d.get('status')
        with conn() as c:
            c.execute('BEGIN IMMEDIATE')
            r = c.execute('SELECT * FROM stock_requests WHERE id=?', (rid,)).fetchone()
            if not r or (u['role'] == 'client' and r['client_id'] != u['id']):
                raise Error('Demande introuvable.', 404)
            if r['status'] != 'En attente':
                raise Error('Cette demande est déjà traitée. Aucun nouveau mouvement n’a été créé.', 409)
            if u['role'] == 'client' and status != 'Annulée':
                raise Error('Seule l’administration peut valider ou refuser cette demande.', 403)
            if status not in ('Validée', 'Refusée', 'Annulée'):
                raise Error('Décision invalide.')
            note = text(d, 'admin_note', required=status == 'Refusée')
            movement_id = None
            if status == 'Validée':
                p = find_product(c, r['product_id'], u)
                delta = r['quantity'] if r['kind'] == 'Entrée' else -r['quantity']
                if p['quantity'] + delta < services['reserved_stock'](c,p['id']):
                    raise Error('Stock insuffisant au moment de la validation. Aucun mouvement effectué.', 409)
                c.execute('UPDATE products SET quantity=quantity+? WHERE id=?', (delta, p['id']))
                cur = c.execute('INSERT INTO movements(product_id,actor_id,delta,note,created_at) VALUES(?,?,?,?,?)',
                                (p['id'], u['id'], delta, f"Demande STK-{rid:04d} · {r['reason']}" + (f' · {note}' if note else ''), now()))
                movement_id = cur.lastrowid
            c.execute('UPDATE stock_requests SET status=?,admin_note=?,processed_by=?,processed_at=?,movement_id=? WHERE id=?',
                      (status, note, u['id'], now(), movement_id, rid))
        return jsonify(ok=True, movement_id=movement_id)

    @app.get('/api/stock/movements')
    @auth('admin', 'client')
    def stock_history():
        u = user(); cond, args = ('1=1', []) if u['role'] == 'admin' else ('p.client_id=?', [u['id']])
        with conn() as c:
            rows = c.execute(f'''SELECT m.*,p.name product,p.reference sku,u.name actor,cl.company company
                FROM movements m JOIN products p ON p.id=m.product_id JOIN users u ON u.id=m.actor_id
                JOIN users cl ON cl.id=p.client_id WHERE {cond} ORDER BY m.id DESC LIMIT 500''', args)
            return jsonify([dict(r) for r in rows])


def filename_display(s):
    return s[:120]
