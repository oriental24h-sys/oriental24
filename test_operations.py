"""Tests for phase 2. All writes use the same isolated DB fixtures as test_app."""
import io, csv, json, unittest, base64
from datetime import datetime, timedelta
from openpyxl import Workbook, load_workbook
from PIL import Image
import test_app as fixtures
system = fixtures.system
from operations import COLUMNS

class OperationsTests(unittest.TestCase):
    setUp = fixtures.PlatformTests.setUp
    data = fixtures.PlatformTests.data
    mutate = fixtures.PlatformTests.mutate

    def rows(self, reference='TEST-001', **kw):
        d=dict(zip(COLUMNS,[reference,'Test destinataire','0600000000','12 avenue Test','Oujda','199,90','Article','Note test']))
        d.update(kw);return d
    def preview(self, rows=None, client=None, fmt='csv', client_id=2):
        rows=rows if rows is not None else [self.rows()]
        if fmt=='csv':
            b=io.StringIO();w=csv.DictWriter(b,fieldnames=COLUMNS,delimiter=';');w.writeheader();w.writerows(rows);content=b.getvalue().encode('utf-8-sig')
        else:
            wb=Workbook();ws=wb.active;ws.title='Colis';ws.append(COLUMNS)
            for row in rows:ws.append([row.get(k) for k in COLUMNS])
            b=io.BytesIO();wb.save(b);content=b.getvalue()
        client=client or self.client
        return client.post('/api/imports/preview',data={'file':(io.BytesIO(content),'test.'+fmt),'client_id':str(client_id)},headers={'X-CSRF-Token':self.data(client)['csrf']})
    def commit(self,b,client=None):return self.mutate(client or self.client,'/api/imports/'+b['id']+'/commit',{})
    def stock_request(self,quantity=5,kind='Entrée',product_id=1):
        return self.mutate(self.client,'/api/stock/requests',{'product_id':product_id,'quantity':quantity,'kind':kind,'reason':'Contrôle stock test'})
    def test_preview_does_not_create_and_confirmation_is_idempotent(self):
        initial=len(self.data(self.client)['parcels'])
        b=self.preview().get_json();self.assertEqual(b['invalid'],0,b)
        self.assertEqual(len(self.data(self.client)['parcels']),initial)
        first=self.commit(b);self.assertEqual(first.status_code,200,first.get_json())
        self.assertEqual(first.get_json()['count'],1)
        again=self.commit(b);self.assertTrue(again.get_json()['already_imported'])
        self.assertEqual(first.get_json()['parcels'],again.get_json()['parcels'])
        self.assertEqual(len(self.data(self.client)['parcels']),initial+1)
        p=next(p for p in self.data(self.client)['parcels'] if p['id']==first.get_json()['parcels'][0]['id'])
        self.assertEqual(p['amount'],199.90);self.assertEqual(p['fee'],25)
        b2=self.preview().get_json();self.assertEqual(b2['invalid'],1)
        self.assertIn('déjà importée',' '.join(b2['rows'][0]['errors']))
    def test_atomic_error_batch_and_report(self):
        b=self.preview([self.rows(),self.rows('TEST-002',telephone='600000000',ville='Ville inconnue',montant='nan')]).get_json()
        self.assertEqual((b['valid'],b['invalid']),(1,1))
        self.assertEqual(self.commit(b).status_code,400)
        report=self.client.get('/api/imports/'+b['id']+'/errors')
        self.assertEqual(report.status_code,200);self.assertIn('TEST-002',report.data.decode('utf-8-sig'))
        self.assertIn('Téléphone',report.data.decode('utf-8-sig'))
    def test_duplicate_reference_within_file_case_insensitive(self):
        b=self.preview([self.rows('ABC'),self.rows('abc')]).get_json()
        self.assertEqual(b['invalid'],1)
        self.assertEqual(self.commit(b).status_code,400)
    def test_import_ownership_and_forced_client(self):
        b=self.preview(client_id=4).get_json();self.assertEqual(b['client'],'Maison Zina')
        self.assertEqual(self.admin.get('/api/imports/'+b['id']).status_code,404)
        self.assertEqual(self.commit(b,self.admin).status_code,404)
        self.assertEqual(self.admin.get('/api/imports/'+b['id']+'/errors').status_code,404)
        self.assertEqual(self.driver.get('/api/imports').status_code,403)
        self.assertEqual(self.driver.get('/api/imports/template.xlsx').status_code,403)
    def test_admin_can_import_for_client(self):
        b=self.preview(client=self.admin,client_id=4).get_json()
        self.assertEqual(b['client'],'Studio Safran')
        r=self.commit(b,self.admin).get_json();pid=r['parcels'][0]['id']
        self.assertEqual(self.client.get('/api/parcels/'+str(pid)).status_code,404)
        self.assertEqual(self.admin.get('/api/parcels/'+str(pid)).get_json()['parcel']['client_id'],4)
    def test_coverage_and_tariffs_rechecked(self):
        b=self.preview().get_json()
        city=next(c for c in self.data(self.admin)['cities'] if c['id']==1)
        self.mutate(self.admin,'/api/cities/1',{**city,'fee':99},'patch')
        self.assertEqual(self.commit(b).status_code,409)
        self.mutate(self.admin,'/api/cities/1',{**city,'delivery':False},'patch')
        self.assertEqual(self.preview().get_json()['invalid'],1)
    def test_expired_import_not_committed(self):
        b=self.preview().get_json()
        with system.conn() as c:c.execute('UPDATE import_batches SET created_at=? WHERE id=?',((datetime.now()-timedelta(hours=1)).isoformat(),b['id']))
        self.assertEqual(self.commit(b).status_code,409)
    def test_parallel_previews_cannot_duplicate_commit(self):
        first=self.preview([self.rows('A'),self.rows('B')]).get_json()
        second=self.preview([self.rows('C'),self.rows('B')]).get_json()
        self.assertEqual(self.commit(first).status_code,200)
        self.assertEqual(self.commit(second).status_code,409)
        with system.conn() as c:
            self.assertIsNone(c.execute("SELECT 1 FROM parcel_external_refs WHERE reference='c'").fetchone())
    def test_xlsx_parsing_formula_numeric_phone(self):
        good=self.preview(fmt='xlsx').get_json();self.assertEqual(good['invalid'],0,good)
        self.assertEqual(self.commit(good).status_code,200)
        bad=self.preview([self.rows('FORMULA',note='=1+1'),self.rows('NUMBER',telephone=600000000)],fmt='xlsx').get_json()
        self.assertEqual(bad['invalid'],2)
    def test_dynamic_template_is_valid_xlsx(self):
        r=self.client.get('/api/imports/template.xlsx');self.assertEqual(r.status_code,200)
        wb=load_workbook(io.BytesIO(r.data));self.assertEqual(wb.sheetnames,['Colis','Villes','Guide'])
        self.assertEqual(wb['Colis']['C2'].value,'0600000000')
        self.assertEqual(wb['Colis']['C2'].number_format,'@')
        self.assertEqual(wb['Villes'].max_row,len([c for c in self.data(self.client)['cities'] if c['delivery']])+1)
        wb.close()
        result=self.client.post('/api/imports/preview',data={'file':(io.BytesIO(r.data),'modele.xlsx')},headers={'X-CSRF-Token':self.data(self.client)['csrf']})
        self.assertEqual(result.status_code,200,result.get_json());self.assertEqual(result.get_json()['total'],1)
    def test_import_row_limit_and_unsupported_format(self):
        r=self.preview([self.rows(str(i)) for i in range(501)])
        self.assertEqual(r.status_code,400)
        r=self.preview([self.rows(str(i)) for i in range(501)],fmt='xlsx')
        self.assertEqual(r.status_code,400)
        r=self.client.post('/api/imports/preview',data={'file':(io.BytesIO(b'bad'),'test.xls')},headers={'X-CSRF-Token':self.data(self.client)['csrf']})
        self.assertEqual(r.status_code,400)
    def test_label_generation_and_access(self):
        own=self.data(self.client)['parcels'][0];other=next(p for p in self.data(self.admin)['parcels'] if p['client_id']!=2)
        r=self.mutate(self.client,'/api/labels',{'ids':[own['id']]});self.assertEqual(r.status_code,200,r.get_json())
        d=r.get_json()[0];image=Image.open(io.BytesIO(base64.b64decode(d['qr'].split(',')[1])));self.assertGreater(image.width,100)
        svg=base64.b64decode(d['barcode'].split(',')[1]).decode();self.assertIn('<svg',svg);self.assertIn(own['tracking'],svg)
        self.assertEqual(self.mutate(self.client,'/api/labels',{'ids':[own['id'],other['id']]}).status_code,404)
        self.assertEqual(self.mutate(self.client,'/api/labels',{'ids':[own['id']]*51}).status_code,400)
        self.assertEqual(self.mutate(self.client,'/api/labels',{'ids':['1 OR 1=1']}).status_code,400)
    def test_scanner_lookup_is_scoped_and_does_not_mutate(self):
        own=self.data(self.driver)['parcels'][0];other=next(p for p in self.data(self.admin)['parcels'] if p['driver_id']!=3)
        r=self.mutate(self.driver,'/api/scan',{'tracking':own['tracking'].lower()});self.assertEqual(r.status_code,200)
        self.assertEqual(r.get_json()['id'],own['id'])
        self.assertEqual(self.mutate(self.driver,'/api/scan',{'tracking':other['tracking']}).status_code,404)
        self.assertEqual(self.mutate(self.driver,'/api/scan',{'tracking':'javascript:alert(1)'}).status_code,400)
        self.assertEqual(self.driver.get('/api/parcels/'+str(own['id'])).get_json()['parcel']['status'],own['status'])
    def test_stock_validation_adds_exactly_one_movement(self):
        r=self.stock_request();self.assertEqual(r.status_code,200)
        rid=r.get_json()['id'];self.assertEqual(self.client.get('/api/products').get_json()[0]['quantity'],24)
        r=self.mutate(self.admin,f'/api/stock/requests/{rid}',{'status':'Validée','admin_note':'Reçu et vérifié'},'patch')
        self.assertEqual(r.status_code,200);self.assertIsNotNone(r.get_json()['movement_id'])
        self.assertEqual(self.client.get('/api/products').get_json()[0]['quantity'],29)
        self.assertEqual(self.mutate(self.admin,f'/api/stock/requests/{rid}',{'status':'Validée'},'patch').status_code,409)
        self.assertEqual(len(self.client.get('/api/stock/movements').get_json()),1)
    def test_stock_reject_cancel_and_role_access(self):
        rid=self.stock_request().get_json()['id']
        self.assertEqual(self.mutate(self.client,f'/api/stock/requests/{rid}',{'status':'Validée'},'patch').status_code,403)
        self.assertEqual(self.mutate(self.admin,f'/api/stock/requests/{rid}',{'status':'Refusée'},'patch').status_code,400)
        self.assertEqual(self.mutate(self.admin,f'/api/stock/requests/{rid}',{'status':'Refusée','admin_note':'Produit non reçu'},'patch').status_code,200)
        rid2=self.stock_request().get_json()['id']
        self.assertEqual(self.mutate(self.client,f'/api/stock/requests/{rid2}',{'status':'Annulée'},'patch').status_code,200)
        self.assertEqual(self.client.get('/api/products').get_json()[0]['quantity'],24)
        self.assertEqual(self.client.get('/api/stock/movements').get_json(),[])
        self.assertEqual(self.driver.get('/api/stock/requests').status_code,403)
    def test_stock_revalidation_prevents_negative_stock(self):
        a=self.stock_request(20,'Sortie').get_json()['id'];b=self.stock_request(20,'Sortie').get_json()['id']
        self.assertEqual(self.mutate(self.admin,f'/api/stock/requests/{a}',{'status':'Validée'},'patch').status_code,200)
        self.assertEqual(self.mutate(self.admin,f'/api/stock/requests/{b}',{'status':'Validée'},'patch').status_code,409)
        self.assertEqual(self.client.get('/api/products').get_json()[0]['quantity'],4)
        self.assertEqual(next(r for r in self.client.get('/api/stock/requests').get_json() if r['id']==b)['status'],'En attente')
    def test_stock_other_client_cannot_read_or_decide(self):
        rid=self.stock_request().get_json()['id']
        other=system.app.test_client();other.post('/api/login',json={'email':'sara@example.test','password':'Oriental24!Demo'})
        self.assertEqual(other.get('/api/stock/requests').get_json(),[])
        self.assertEqual(self.mutate(other,f'/api/stock/requests/{rid}',{'status':'Annulée'},'patch').status_code,404)
        self.assertEqual(self.mutate(other,'/api/stock/requests',{'product_id':1,'quantity':1,'kind':'Entrée','reason':'Test'}).status_code,404)

if __name__=='__main__':unittest.main(verbosity=2)
